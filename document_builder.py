from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

def create_automotive_report():
    doc = Document()
    
    # تنسيق الخط العام
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Arial'
    font.size = Pt(12)

    # إضافة العنوان الرئيسي
    title = doc.add_heading('Automotive Strategic Analysis Report', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # القسم الأول: المقدمة
    doc.add_heading('1. Executive Summary', level=1)
    doc.add_paragraph(
        "This strategic report analyzes vehicle sales performance (ICE vs EV) "
        "during periods of economic instability from 2019 to 2023."
    )

    # القسم الثاني: المنهجية
    doc.add_heading('2. Methodology', level=1)
    doc.add_paragraph(
        "The analysis utilizes synthetic monthly sales data to evaluate the "
        "resilience of different vehicle segments during identified crisis years."
    )

    # القسم الثالث: التحليل المرئي
    doc.add_heading('3. Data Visualizations', level=1)
    
    # إضافة الرسم الأول: اتجاهات المبيعات
    doc.add_heading('3.1 Sales Growth Trends', level=2)
    if os.path.exists('figures/sales_trend.png'):
        doc.add_picture('figures/sales_trend.png', width=Inches(5.5))
        doc.add_paragraph("Figure 1: Comparison of monthly sales trends for ICE and EV models.")

    doc.add_page_break()

    # إضافة الرسم الثاني: تأثير الأزمات
    doc.add_heading('3.2 Economic Crisis Impact', level=2)
    if os.path.exists('figures/crisis_impact.png'):
        doc.add_picture('figures/crisis_impact.png', width=Inches(5.5))
        doc.add_paragraph("Figure 2: Performance breakdown during normal vs. crisis economic periods.")

    # القسم الرابع: التوصيات
    doc.add_heading('4. Strategic Recommendations', level=1)
    doc.add_paragraph(
        "Based on the observed resilience of Electric Vehicles, strategic focus should "
        "shift towards sustainable automotive technologies to mitigate economic risks."
    )

    output_path = 'outputs/Automotive_Strategic_Report.docx'
    doc.save(output_path)
    print(f"Word document generated successfully at: {output_path}")

if __name__ == "__main__":
    create_automotive_report()